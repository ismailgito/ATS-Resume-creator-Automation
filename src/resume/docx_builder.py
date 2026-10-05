"""ATS-Safe DOCX Resume Document Builder."""

from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from src.models.profile import CandidateProfile
from src.models.resume import TailoredResume


def build_docx_resume(candidate: CandidateProfile, tailored: TailoredResume, output_path: Path) -> Path:
    """Generate ATS-optimized Word document (.docx)."""
    doc = Document()

    # Set 0.75 - 1.0 inch standard margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Candidate Name (Header)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(candidate.full_name)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(26, 36, 43)

    # Subtitle / Headline
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(candidate.headline)
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(70, 80, 95)

    # Contact Details
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_text = (
        f"{candidate.contact.email}  |  {candidate.contact.phone}  |  {candidate.contact.location}\n"
        f"LinkedIn: {candidate.contact.linkedin}  |  Portfolio: {candidate.contact.portfolio}"
    )
    contact_run = contact_p.add_run(contact_text)
    contact_run.font.name = "Calibri"
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = RGBColor(80, 80, 80)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 44, 87)

    # 1. Professional Summary
    add_section_header("Professional Summary")
    sum_p = doc.add_paragraph()
    sum_p.paragraph_format.space_after = Pt(6)
    sum_run = sum_p.add_run(tailored.summary)
    sum_run.font.name = "Calibri"
    sum_run.font.size = Pt(10)

    # 2. Core Competencies & Skills
    add_section_header("Core Competencies & Skills")
    skills_p = doc.add_paragraph()
    skills_p.paragraph_format.space_after = Pt(4)
    skills_run = skills_p.add_run(" • ".join(tailored.skills))
    skills_run.font.name = "Calibri"
    skills_run.font.size = Pt(9.5)

    tools_p = doc.add_paragraph()
    tools_p.paragraph_format.space_after = Pt(6)
    tools_label = tools_p.add_run("Platforms & Tools: ")
    tools_label.bold = True
    tools_label.font.name = "Calibri"
    tools_label.font.size = Pt(9.5)
    tools_val = tools_p.add_run(", ".join(candidate.tools))
    tools_val.font.name = "Calibri"
    tools_val.font.size = Pt(9.5)

    # 3. Professional Experience
    add_section_header("Professional Experience")
    for role in tailored.experience:
        role_p = doc.add_paragraph()
        role_p.paragraph_format.space_before = Pt(4)
        role_p.paragraph_format.space_after = Pt(1)

        title_r = role_p.add_run(f"{role.title} — {role.company}")
        title_r.font.name = "Calibri"
        title_r.font.size = Pt(10.5)
        title_r.font.bold = True

        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_after = Pt(2)
        meta_r = meta_p.add_run(f"{role.location}  |  {role.period}")
        meta_r.font.name = "Calibri"
        meta_r.font.size = Pt(9)
        meta_r.font.italic = True
        meta_r.font.color.rgb = RGBColor(90, 90, 90)

        for bullet in role.bullets:
            b_p = doc.add_paragraph(style="List Bullet")
            b_p.paragraph_format.space_after = Pt(2)
            b_run = b_p.add_run(bullet)
            b_run.font.name = "Calibri"
            b_run.font.size = Pt(9.5)

    # 4. Key Achievements & Impact
    add_section_header("Key Achievements & Track Record")
    for ach in candidate.achievements:
        ach_p = doc.add_paragraph(style="List Bullet")
        ach_p.paragraph_format.space_after = Pt(2)
        m_run = ach_p.add_run(f"{ach.metric}: ")
        m_run.bold = True
        m_run.font.name = "Calibri"
        m_run.font.size = Pt(9.5)
        s_run = ach_p.add_run(f"{ach.statement} ({ach.time_period})")
        s_run.font.name = "Calibri"
        s_run.font.size = Pt(9.5)

    # 5. Education & Certifications
    add_section_header("Education & Certifications")
    for edu in candidate.education:
        edu_p = doc.add_paragraph(style="List Bullet")
        edu_p.paragraph_format.space_after = Pt(1)
        edu_run = edu_p.add_run(f"{edu.degree}, {edu.institution} ({edu.start_year} - {edu.end_year})")
        edu_run.font.name = "Calibri"
        edu_run.font.size = Pt(9.5)

    for cert in candidate.certifications:
        cert_p = doc.add_paragraph(style="List Bullet")
        cert_p.paragraph_format.space_after = Pt(1)
        c_run = cert_p.add_run(f"{cert.name} — {cert.issuer} ({cert.date})")
        c_run.font.name = "Calibri"
        c_run.font.size = Pt(9.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
